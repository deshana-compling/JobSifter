import re
import phonenumbers
from pathlib import Path
from pypdf import PdfReader
from docx import Document

#Function to extract text from pdf
def pdfextractor(file):
    
    ''' EXTRACT TEXT FROM A PDF. CV'''

    reader = PdfReader(file)
    text = ''
    for page in reader.pages:
        pagetext = page.extract_text()
        if pagetext:
            text += pagetext + '\n'
    return text 

#Function to extract text from doc
def docextractor(file):

    ''' EXTRACT TEXT FROM A DOCX. CV'''

    doc = Document(file)
    text = ''
    for para in doc.paragraphs:
        text += para.text + '\n'
    
    for table in doc.tables:  #Edge Case: Sometimes CV contains tables!
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + '\n'
    
    for section in doc.sections:  #Headers and footers
        for para in section.header.paragraphs:
            text += para.text + '\n'
        for para in section.footer.paragraphs:
            text += para.text + '\n'
    return text


def redactor(text):

    '''REMOVE PERSONAL DETAILS & UNWANTED ELEMENTS FROM CV'''

    #Email
    text = re.sub(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b", "[EMAIL REDACTED]", text)

    #URLs/Links
    text = re.sub(r"https?://\S+|www\.\S+", "[URL REDACTED]", text)
    text = re.sub(r"\b(?:linkedin\.com|github\.com|twitter\.com|x\.com)/\S+", "[URL REDACTED]", text, flags=re.IGNORECASE)

    #Phone numbers
    phones = list(phonenumbers.PhoneNumberMatcher(text, None))
    for phone in reversed(phones):
        text = (text[:phone.start] + "[PHONE REDACTED]" + text[phone.end:])
    
    #Replace multiple spaces/tabs with single space
    text = re.sub(r"[ \t]+", " ", text)

    #Remove spaces at the beginning/end of lines
    text = re.sub(r" *\n *", "\n", text)

    #Remove excessive blank lines
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text

#FINAL Preprocessing function
def processor(file):
    
    file_ext = Path(file.name).suffix.lower()
 
    if file_ext == '.pdf':
        text = pdfextractor(file)
    elif file_ext == '.docx':
        text = docextractor(file)
    elif file_ext == '.doc':
        raise ValueError(".doc files are not supported. Please convert to .docx first (e.g. via LibreOffice or Word's 'Save As'). OR upload .pdf")
    else:
        raise ValueError('Unsupported file type. Please upload a PDF or DOCX file.')
    if not text.strip():
        raise ValueError("No extractable text found. The file may be a scanned/image-based PDF,"
            "which would need OCR (e.g. pytesseract) before this pipeline can process it.")
 
    return redactor(text)